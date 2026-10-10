import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)  # 16:9 Widescreen format
prs.slide_height = Inches(7.5)

# --- COLOR PALETTE (Cybersecurity Charcoal Theme) ---
COLOR_BG = RGBColor(18, 20, 23)        # Graphite Black (#121417)
COLOR_CARD = RGBColor(26, 29, 36)      # Dark Slate (#1A1D24)
COLOR_PRIMARY = RGBColor(16, 185, 129) # Cyber Emerald Green (#10B981)
COLOR_TEXT_MAIN = RGBColor(243, 244, 246) # Off-white (#F3F4F6)
COLOR_TEXT_MUTED = RGBColor(156, 163, 175) # Muted Gray (#9CA3AF)
COLOR_ACCENT = RGBColor(56, 189, 248)  # Cyan (#38BDF8)

def apply_background(slide):
    """Adds a dark charcoal background to a slide."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.fill.background()

def create_header(slide, title_text, category="CYBERGUARD PLATFORM"):
    """Creates a consistent, non-overlapping header zone."""
    # Category Label
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
    tf_cat = cat_box.text_frame
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_PRIMARY

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(26)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT_MAIN

# ==============================================================================
# SLIDE 1: Title Slide
# ==============================================================================
slide_layout = prs.slide_layouts[6] # Blank
slide1 = prs.slides.add_slide(slide_layout)
apply_background(slide1)

# Large Accent Card
card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5))
card.fill.solid()
card.fill.fore_color.rgb = COLOR_CARD
card.line.color.rgb = COLOR_PRIMARY

# Title Text Box
tbox = slide1.shapes.add_textbox(Inches(2.0), Inches(2.0), Inches(9.333), Inches(3.5))
tf = tbox.text_frame
tf.word_wrap = True

p0 = tf.paragraphs[0]
p0.text = "CYBERGUARD"
p0.font.size = Pt(44)
p0.font.bold = True
p0.font.color.rgb = COLOR_PRIMARY

p1 = tf.add_paragraph()
p1.text = "AI-Powered Threat, Phishing & Impersonation Detection System"
p1.font.size = Pt(20)
p1.font.color.rgb = COLOR_TEXT_MAIN
p1.space_after = Pt(20)

p2 = tf.add_paragraph()
p2.text = "Domain: Cybersecurity + Artificial Intelligence\nPresenter : Sarojini Patra & Team\nInstitution: Suddhananda Engineering & Research Centre"
p2.font.size = Pt(14)
p2.font.color.rgb = COLOR_TEXT_MUTED

# ==============================================================================
# SLIDE 2: Problem Statement & Background
# ==============================================================================
slide2 = prs.slides.add_slide(slide_layout)
apply_background(slide2)
create_header(slide2, "Problem Statement & Cybersecurity Challenges")

# Box 1: The Threat
b1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
b1.fill.solid()
b1.fill.fore_color.rgb = COLOR_CARD
b1.line.color.rgb = COLOR_ACCENT

tf1 = b1.text_frame
tf1.margin_left = Inches(0.3)
tf1.margin_right = Inches(0.3)
tf1.word_wrap = True
p = tf1.paragraphs[0]
p.text = "The Evolving Cyber Threat Landscape"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT

bullets1 = [
    
    "Rapid adoption of digital payments, cloud services, and e-governance has expanded the attack surface.",
    "Attackers utilize Generative AI for phishing emails, fake websites, cloned voices, and deepfake media.",
    "Human-targeted social engineering bypasses legacy technical firewalls."
]
for b in bullets1:
    p = tf1.add_paragraph()
    p.text = "• " + b
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_after = Pt(10)

# Box 2: The Gap & Solution
b2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
b2.fill.solid()
b2.fill.fore_color.rgb = COLOR_CARD
b2.line.color.rgb = COLOR_PRIMARY

tf2 = b2.text_frame
tf2.margin_left = Inches(0.3)
tf2.margin_right = Inches(0.3)
tf2.word_wrap = True
p = tf2.paragraphs[0]
p.text = "Why Rule-Based Systems Fail"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY

bullets2 = [
    "Traditional security rules miss zero-day AI threats and dynamic URL structures.",
    "Lack of Explainable AI (XAI) leaves security analysts without key decision rationale.",
    "CYBERGUARD Solution: An end-to-end AI platform delivering real-time threat detection, risk scoring, and automated mitigation response."
]
for b in bullets2:
    p = tf2.add_paragraph()
    p.text = "• " + b
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_after = Pt(10)

# ==============================================================================
# SLIDE 3: Key Threat Scenarios (4-Grid Card Display)
# ==============================================================================
slide3 = prs.slides.add_slide(slide_layout)
apply_background(slide3)
create_header(slide3, "Core Threat Modules & Scenarios")

scenarios = [
    ("1. Phishing & URL Detection", "Scans emails, SMS, and links for domain spoofing, urgency keywords, raw IP patterns, and high-risk TLDs.", COLOR_PRIMARY),
    ("2. Call & Digital Impersonation", "Checks caller IDs against blacklists, 140-series telemarketer patterns, and transcript fraud triggers.", COLOR_ACCENT),
    ("3. Account Takeover (ATO)", "Detects brute-force login sequences, abnormal geographic locations, and unrecognized device signatures.", COLOR_PRIMARY),
    ("4. Multimodal Deepfakes", "Inspects images and audio files for missing EXIF metadata, sample rate anomalies, and voice cloning signatures.", COLOR_ACCENT)
]

coords = [
    (Inches(0.8), Inches(1.8)),
    (Inches(6.8), Inches(1.8)),
    (Inches(0.8), Inches(4.5)),
    (Inches(6.8), Inches(4.5))
]

for idx, (title, desc, border_col) in enumerate(scenarios):
    x, y = coords[idx]
    box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.7), Inches(2.3))
    box.fill.solid()
    box.fill.fore_color.rgb = COLOR_CARD
    box.line.color.rgb = border_col
    
    tf = box.text_frame
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = border_col
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_TEXT_MAIN

# ==============================================================================
# SLIDE 4: Architecture & Workflow
# ==============================================================================
slide4 = prs.slides.add_slide(slide_layout)
apply_background(slide4)
create_header(slide4, "System Architecture & Processing Pipeline")

# Workflow Pipeline Diagram
steps = [
    "1. Multi-Source Ingestion\n(Text, URLs, Calls, Audio, Images)",
    "2. AI Detection Engine\n(Regex, Heuristics, Metadata Analysis)",
    "3. Risk Scoring Engine\n(Capped 0-100 Score & Severity Mapping)",
    "4. Explainable AI & Action\n(Reasoning Generation + Automated Response)"
]

for idx, step in enumerate(steps):
    x = Inches(0.8 + idx * 3.0)
    box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(2.7), Inches(1.8))
    box.fill.solid()
    box.fill.fore_color.rgb = COLOR_CARD
    box.line.color.rgb = COLOR_PRIMARY
    
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = step
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.alignment = PP_ALIGN.CENTER

# Tech Stack Container
stack_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.0), Inches(11.7), Inches(2.6))
stack_box.fill.solid()
stack_box.fill.fore_color.rgb = COLOR_CARD
stack_box.line.color.rgb = COLOR_ACCENT

tf = stack_box.text_frame
tf.margin_left = Inches(0.3)
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Technology Stack"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT

bullets = [
    "Backend & Processing: Python 3, Flask, SQLite3, Mutagen (Audio Metadata), Pillow (Image EXIF Processing)",
    "Frontend & Dashboard: HTML5, CSS3 (SOC Dark Theme), JavaScript, Chart.js for real-time analytics",
    "Security Implementation: Werkzeug password hashing, session revocation, local timezone conversion"
]
for b in bullets:
    p = tf.add_paragraph()
    p.text = "• " + b
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_after = Pt(8)

# ==============================================================================
# SLIDE 5: Risk Scoring & Explainable AI (XAI)
# ==============================================================================
slide5 = prs.slides.add_slide(slide_layout)
apply_background(slide5)
create_header(slide5, "Threat Risk Scoring & Explainable AI (XAI)")

box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
box.fill.solid()
box.fill.fore_color.rgb = COLOR_CARD
box.line.color.rgb = COLOR_PRIMARY

tf = box.text_frame
tf.margin_left = Inches(0.4)
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Transparent Risk Categorization"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY

items = [
    ("Risk Scale:", "Safe (0-19) -> Medium (20-39) -> High (40-69) -> Critical (70-100)"),
    ("XAI Philosophy:", "Instead of opaque output like 'Threat Detected', CYBERGUARD generates detailed human-readable evidence detailing exact risk factors."),
    ("Phishing XAI Example:", "Critical Risk | Contains high-pressure urgency language ('verify now') and embedded URL redirects using raw IP addresses."),
    ("Voice Clone XAI Example:", "Critical Risk | Audio sample rate (16,000 Hz) and single-channel signature match standard AI speech synthesis engines.")
]

for title, desc in items:
    p = tf.add_paragraph()
    p.text = title + " "
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_ACCENT
    
    p.text += desc
    p.font.bold = False
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_after = Pt(12)

# ==============================================================================
# SLIDE 6: Automated Incident Response
# ==============================================================================
slide6 = prs.slides.add_slide(slide_layout)
apply_background(slide6)
create_header(slide6, "Automated Mitigation & Incident Response")

box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
box.fill.solid()
box.fill.fore_color.rgb = COLOR_CARD
box.line.color.rgb = COLOR_ACCENT

tf = box.text_frame
tf.margin_left = Inches(0.4)
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Actionable Playbooks by Threat Level"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT

responses = [
    ("Phishing & Malicious URLs:", "Quarantine email message, block embedded domain links, display warning banner, and notify SOC."),
    ("Vishing & Call Fraud:", "Block incoming caller ID immediately, restrict audio source trust, and trigger automated fraud alert."),
    ("Account Takeover Anomaly:", "Revoke active login session, temporarily lock account, and mandate Multi-Factor Authentication (MFA)."),
    ("Multimedia Deepfakes:", "Apply digital authenticity warning watermark, restrict playback, and flag for manual analyst verification.")
]

for title, desc in responses:
    p = tf.add_paragraph()
    p.text = "• " + title + " "
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_PRIMARY
    
    p.text += desc
    p.font.bold = False
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_after = Pt(14)

# ==============================================================================
# SLIDE 7: Command Dashboard & Analytics
# ==============================================================================
slide7 = prs.slides.add_slide(slide_layout)
apply_background(slide7)
create_header(slide7, "Cybersecurity Command Dashboard")

# Left Column
b1 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
b1.fill.solid()
b1.fill.fore_color.rgb = COLOR_CARD
b1.line.color.rgb = COLOR_PRIMARY

tf1 = b1.text_frame
tf1.margin_left = Inches(0.3)
tf1.word_wrap = True
p = tf1.paragraphs[0]
p.text = "Real-Time SOC Analytics"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY

features1 = [
    "Event Counters: Tracks Total Scans, High/Critical Threat counts in real time.",
    "Interactive Visualizations: Chart.js bar graphs dynamically display threat distribution across Phishing, Calls, ATO, Images, and Audio.",
    "Local Timezone Auditing: Auto-converts SQLite UTC timestamps to local system time."
]
for f in features1:
    p = tf1.add_paragraph()
    p.text = "• " + f
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_after = Pt(12)

# Right Column
b2 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
b2.fill.solid()
b2.fill.fore_color.rgb = COLOR_CARD
b2.line.color.rgb = COLOR_ACCENT

tf2 = b2.text_frame
tf2.margin_left = Inches(0.3)
tf2.word_wrap = True
p = tf2.paragraphs[0]
p.text = "Live Audit Trail & History"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT

features2 = [
    "Timestamped Incident Log: Maintains a historical record of all evaluated events.",
    "Severity Badges: Visually differentiates Safe, Medium, High, and Critical risk events.",
    "Full Transparency: Displays input payloads alongside generated XAI explanations and mitigation recommendations."
]
for f in features2:
    p = tf2.add_paragraph()
    p.text = "• " + f
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_after = Pt(12)

# ==============================================================================
# SLIDE 8: Live Demonstration Guide
# ==============================================================================
slide8 = prs.slides.add_slide(slide_layout)
apply_background(slide8)
create_header(slide8, "Live Demonstration Walkthrough")

box = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
box.fill.solid()
box.fill.fore_color.rgb = COLOR_CARD
box.line.color.rgb = COLOR_PRIMARY

tf = box.text_frame
tf.margin_left = Inches(0.4)
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Demonstration Steps for Judges"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY

steps_demo = [
    ("Step 1: Phishing Module", "Input 'URGENT: Your account access is suspended. Verify now at http://login-secure.xyz' -> Show risk score & XAI reasoning."),
    ("Step 2: Call Impersonation", "Input '+919876543210' with OTP transcript -> Show vishing alert."),
    ("Step 3: Account Takeover", "Set 6 failed attempts + unusual location -> Show automated session revocation trigger."),
    ("Step 4: Audio Deepfake", "Upload audio voice recording -> Inspect sample rate metadata extraction."),
    ("Step 5: Command Dashboard", "Refresh dashboard to verify live Chart.js update and audit log entries.")
]

for title, desc in steps_demo:
    p = tf.add_paragraph()
    p.text = title + ": "
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_ACCENT
    
    p.text += desc
    p.font.bold = False
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_after = Pt(10)

# ==============================================================================
# SLIDE 9: Innovation & Future Roadmap
# ==============================================================================
slide9 = prs.slides.add_slide(slide_layout)
apply_background(slide9)
create_header(slide9, "Innovation & Future Enhancements")

box = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
box.fill.solid()
box.fill.fore_color.rgb = COLOR_CARD
box.line.color.rgb = COLOR_ACCENT

tf = box.text_frame
tf.margin_left = Inches(0.4)
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Future Scalability Features"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT

innovations = [
    ("MITRE ATT&CK Mapping:", "Tagging all threat classifications to official technique IDs (e.g., T1566 Phishing, T1110 Brute Force)."),
    ("Network & API Log Scanner:", "Expanding threat analysis to raw server logs for DDoS detection and API rate-limit abuse."),
    ("One-Click SOC Playbooks:", "Adding dynamic action triggers on the dashboard (e.g., [Block IP], [Export Audit PDF])."),
    ("Privacy-Preserving AI:", "Integrating federated learning models for secure enterprise and university deployment.")
]

for title, desc in innovations:
    p = tf.add_paragraph()
    p.text = "• " + title + " "
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_PRIMARY
    
    p.text += desc
    p.font.bold = False
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_after = Pt(14)

# ==============================================================================
# SLIDE 10: Conclusion
# ==============================================================================
slide10 = prs.slides.add_slide(slide_layout)
apply_background(slide10)

# Conclusion Card
card = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5))
card.fill.solid()
card.fill.fore_color.rgb = COLOR_CARD
card.line.color.rgb = COLOR_PRIMARY

tbox = slide10.shapes.add_textbox(Inches(2.0), Inches(2.0), Inches(9.333), Inches(3.5))
tf = tbox.text_frame
tf.word_wrap = True

p0 = tf.paragraphs[0]
p0.text = "CYBERGUARD"
p0.font.size = Pt(36)
p0.font.bold = True
p0.font.color.rgb = COLOR_PRIMARY

p1 = tf.add_paragraph()
p1.text = "Intelligent AI Cyber Defence & Digital Trust Platform"
p1.font.size = Pt(20)
p1.font.color.rgb = COLOR_TEXT_MAIN
p1.space_after = Pt(20)

p2 = tf.add_paragraph()
p2.text = "Protecting both the Technology Layer and the Human Layer of Cybersecurity.\n\nThank You! Questions & Discussion."
p2.font.size = Pt(16)
p2.font.color.rgb = COLOR_TEXT_MUTED

# Save presentation
output_path = "CYBERGUARD_Hackathon_Presentation.pptx"
prs.save(output_path)
print(f"SUCCESS: Presentation saved cleanly as '{output_path}'")